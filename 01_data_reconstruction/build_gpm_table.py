# -*- coding: utf-8 -*-
"""build_gpm_table.py — 将 GPM 多时效独立验证写入 R3 稿 §3j：
1) 在 Table 16 题注后插入 GPM 验证正文段
2) 插入 Table 17 题注 + 表格（复制 Table 15(索引15) 格式，8 列）
3) 更新 Fig.12 题注 (b) 面板加入 GPM
"""
import sys, copy
sys.path.insert(0, r'C:\Users\yg181\AppData\Roaming\Python\Python313\site-packages')
import docx
from docx.shared import Pt

DOC = r'C:\Users\yg181\Desktop\论文三\Manuscript_R3_WAF_draft.docx'
d = docx.Document(DOC)
ps = d.paragraphs

# 1. 定位 Table 16 题注段（语义定位）
cap16 = None
for p in ps:
    if p.text.strip().startswith('Table 16. CHM independent verification at accumulated leads'):
        cap16 = p
        break
assert cap16 is not None, 'Table 16 题注未找到'

# 2. 定位 Fig.12 题注段
fig12 = None
for p in ps:
    if p.text.strip().startswith('Fig. 12. Accumulation-scale dependence'):
        fig12 = p
        break
assert fig12 is not None

# 3. GPM 正文段
body_text = ('Independent satellite verification against GPM IMERG (V07B; 30-min rates summed over the '
             'identical accumulation windows; init 00Z only; 638 / 636 / 634 samples; 828 of 925 grid points '
             'covered) confirms the same reversal with a fully independent, high-frequency reference (Table 17, '
             'Fig. 12b). Against GPM the DL corrections improve RMSE by +14.0% / +15.5% / +13.4% (APCNet) and '
             '+18.4% / +16.7% / +18.1% (U-Net) at 24 / 72 / 120 h, again exceeding every statistical baseline '
             '(QM +3.0% / -0.1% / +1.3%; BinCM +9.0% / +5.5% / +2.6%; OLS +7.6% / +6.4% / +4.0%), and at 24 h '
             'both networks also exceed the ERA5 reference itself (+7.8%), whose own RMSE skill grows to +23.4% '
             'and +39.3% at 72-120 h as initial-condition error accumulates. The GPM numbers therefore reproduce '
             'the CHM gauge result at every lead and remove the 12-h vs 24-h daily-coverage caveat of the '
             'gauge comparison. Together with the 3-h GPM result (APCNet -29.3% against GPM at the native 3-h '
             'scale, Section 3f), the two independent products bound the scale dependence from both sides: the '
             'same network, trained on the same ERA5 target, loses skill against the observation at 3 h and '
             'gains it at all operational accumulated scales.')

cap17_text = ('Table 17. GPM IMERG independent verification at accumulated leads (init 00Z only; 30-min IMERG '
              'rates aggregated over [init, init+fhr]; 638 / 636 / 634 samples; domain mean over the 828 '
              'GPM-covered grid points; mm per lead). Delta RMSE is the RMSE improvement relative to raw GFS '
              '(positive = improvement). QM / BinCM / OLS are the grid-point baselines of Section 2c, calibrated '
              'on the training period only.')

# 4. 复制 Table 15（索引15）格式并改数据
src_tbl = d.tables[15]
new_tbl = copy.deepcopy(src_tbl._tbl)
# 新表数据（8 列，3 数据行 + 表头）
gpm_rows = [
    ['Lead', 'Samples', 'GFS RMSE (mm)', 'APCNet dRMSE (%)', 'U-Net dRMSE (%)',
     'QM dRMSE (%)', 'BinCM dRMSE (%)', 'OLS dRMSE (%)'],
    ['24 h', '638', '5.30', '+14.0', '+18.4', '+3.0', '+9.0', '+7.6'],
    ['72 h', '636', '10.23', '+15.5', '+16.7', '-0.1', '+5.5', '+6.4'],
    ['120 h', '634', '14.85', '+13.4', '+18.1', '+1.3', '+2.6', '+4.0'],
]
from docx.oxml.ns import qn
tbl_obj = None
# 用 lxml 操作新表：清空行 -> 重建
import lxml.etree as etree
# 移除所有 tr 行，重建
for tr in new_tbl.findall(qn('w:tr')):
    new_tbl.remove(tr)
# 从 src 表复制表头行格式：取 src 第一行作模板
src_trs = src_tbl._tbl.findall(qn('w:tr'))
tpl_header = src_trs[0]
tpl_data = src_trs[1]

def make_tr(tpl, texts):
    tr = copy.deepcopy(tpl)
    tcs = tr.findall(qn('w:tc'))
    assert len(tcs) >= len(texts)
    for tc, txt in zip(tcs[:len(texts)], texts):
        # 清空 tc 内段落文字（保留首段格式）
        paras = tc.findall(qn('w:p'))
        for p in paras[1:]:
            tc.remove(p)
        p0 = paras[0]
        # 移除 runs
        for r in p0.findall(qn('w:r')):
            p0.remove(r)
        # 新建 run 并设文本
        import docx.oxml as oxml
        run = oxml.shared.OxmlElement('w:r')
        rPr = oxml.shared.OxmlElement('w:rPr')
        rFonts = oxml.shared.OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rPr.append(rFonts)
        sz = oxml.shared.OxmlElement('w:sz')
        sz.set(qn('w:val'), '20')  # 10pt
        rPr.append(sz)
        run.append(rPr)
        t = oxml.shared.OxmlElement('w:t')
        t.text = txt
        t.set(qn('xml:space'), 'preserve')
        run.append(t)
        p0.append(run)
    return tr

new_tbl.append(make_tr(tpl_header, gpm_rows[0]))
for row in gpm_rows[1:]:
    new_tbl.append(make_tr(tpl_data, row))

# 5. 插入：先插表格，再在表格前插题注段，再在题注前插正文段（倒序 addnext）
cap16._p.addnext(new_tbl)
# 题注段
p_cap = d.add_paragraph()
r_cap = p_cap.add_run(cap17_text)
r_cap.font.name = 'Times New Roman'
r_cap.font.size = Pt(10)
cap16._p.addnext(p_cap._p)   # 插在表格前
# 正文段
p_body = d.add_paragraph()
r_body = p_body.add_run(body_text)
r_body.font.name = 'Times New Roman'
r_body.font.size = Pt(10)
cap16._p.addnext(p_body._p)  # 插在题注前

# 6. 更新 Fig.12 题注 (b) 面板
new_fig = fig12.text.replace(
    '(b) RMSE improvement relative to GFS in the CHM independent verification at 24- / 72- / 120-h leads',
    '(b) RMSE improvement relative to GFS in the CHM gauge and GPM IMERG independent verifications at 24- / '
    '72- / 120-h leads')
for r in list(fig12.runs):
    r.text = ''
fr = fig12.add_run(new_fig)
fr.font.name = 'Times New Roman'
fr.font.size = Pt(10)

d.save(DOC)
print('✅ 已写入: GPM 正文段 + Table 17 题注/表格 + Fig.12 题注更新')
